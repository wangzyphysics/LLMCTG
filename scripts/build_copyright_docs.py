from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
DELIVERY = ROOT / "deliverables"
SOFTWARE_NAME = "LLM中文文本生成系统"
VERSION = "V1.0"


def set_font(run, east_asia="宋体", latin="Arial", size=10.5, bold=None):
    run.font.name = latin
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), latin)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), latin)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margin(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def configure_document(doc, footer_text=SOFTWARE_NAME):
    section = doc.sections[0]
    section.top_margin = Cm(2.4)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.6)
    section.right_margin = Cm(2.4)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)
    for name, size in (("Title", 24), ("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11)):
        style = styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
    title_ppr = styles["Title"]._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run(footer_text + "  ·  ")
    set_font(run, size=8)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)


def add_title_page(doc, subtitle):
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(100)
    run = p.add_run(SOFTWARE_NAME)
    set_font(run, "黑体", "Arial", 24, True)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_before = Pt(18)
    run = p2.add_run(subtitle)
    set_font(run, "黑体", "Arial", 18, True)
    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.paragraph_format.space_before = Pt(26)
    run = p3.add_run(VERSION)
    set_font(run, "宋体", "Arial", 14)
    p4 = doc.add_paragraph()
    p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p4.paragraph_format.space_before = Pt(190)
    run = p4.add_run("编制日期：2026 年 9 月")
    set_font(run, "宋体", "Arial", 11)
    doc.add_page_break()


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.add_run(text)
    return p


def add_body(doc, text, bold_lead=""):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0.74)
    if bold_lead:
        r = p.add_run(bold_lead)
        set_font(r, "宋体", "Arial", 10.5, True)
    r = p.add_run(text)
    set_font(r)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Cm(0.8)
        r = p.add_run(item)
        set_font(r)


def add_code(doc, text):
    for line in text.strip().splitlines():
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.8)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(line)
        set_font(r, "等线", "Consolas", 9)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, "1F4E78")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(str(header))
        set_font(r, "黑体", "Arial", 9.5, True)
        r.font.color.rgb = RGBColor(255, 255, 255)
        set_cell_margin(cell)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            if row_index % 2:
                set_cell_shading(cells[index], "F2F6FA")
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[index].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if index == 0 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(value))
            set_font(r, size=9)
            set_cell_margin(cells[index])
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Cm(width)
    doc.add_paragraph()
    return table


def build_manual():
    doc = Document()
    configure_document(doc)
    add_title_page(doc, "用户操作手册")
    add_heading(doc, "文档说明", 1)
    add_body(doc, "本手册面向教师、编辑、研究人员和系统维护人员，说明软件的安装、配置、内容生成、质量评估、批量处理、模型微调和结果查看方法。基础模式无需网络、显卡或外部模型即可完成全部工作流演示。")
    add_table(doc, ["项目", "说明"], [
        ["软件名称", SOFTWARE_NAME], ["版本", VERSION], ["运行方式", "命令行应用"],
        ["基础环境", "Python 3.10 及以上"], ["默认模式", "离线模板生成与规则质量评估"],
        ["可选能力", "本地 Transformers 模型生成与 LoRA 微调"],
    ], [3.2, 12.5])
    add_heading(doc, "目录", 1)
    add_bullets(doc, [
        "1 软件概述", "2 安装与启动", "3 快速演示", "4 单篇内容生成", "5 文本质量评估",
        "6 批量任务处理", "7 训练数据准备", "8 LoRA 模型微调", "9 配置说明", "10 输出文件说明",
        "11 常见问题与故障处理", "12 数据安全与使用边界", "13 软件卸载",
    ])
    add_heading(doc, "1 软件概述", 1)
    add_heading(doc, "1.1 软件用途", 2)
    add_body(doc, "系统围绕小学、初中和高中三个阅读等级，完成中文阅读材料生成、自动质量检查和批量报告导出。它解决原始研究脚本入口分散、路径依赖、环境配置不一致和结果难以汇总的问题，把模型生成、文本统计和分级评估组织为可重复执行的软件流程。")
    add_heading(doc, "1.2 核心功能", 2)
    add_bullets(doc, [
        "按主题、学段、关键词和附加要求生成中文阅读材料。",
        "使用离线模板后端完成安装验收，或接入用户拥有的本地语言模型。",
        "计算汉字数、句数、段落数、平均句长、重复率、英文比例和难度匹配度。",
        "按优秀、良好、合格、待改进给出综合结论，并列出警告和修改建议。",
        "从 JSONL 读取批量任务，隔离单条异常并输出 HTML、CSV、JSON 和 JSONL 报告。",
        "把带学段标签的文本转换为训练数据，并对兼容模型执行 LoRA 微调。",
    ])
    add_heading(doc, "1.3 工作流程", 2)
    add_body(doc, "用户先选择生成后端并给出主题。系统校验请求后生成正文，将正文交给评估器计算统计指标和综合分，再把文章、分数、提示与元数据写入结构化结果。批量模式会重复这一流程，并在结束后生成汇总报告。")
    add_heading(doc, "2 安装与启动", 1)
    add_heading(doc, "2.1 环境要求", 2)
    add_table(doc, ["类别", "最低要求", "建议配置"], [
        ["操作系统", "Windows 10、常见 Linux 或 macOS", "64 位桌面或服务器系统"],
        ["Python", "3.10", "3.11 或 3.12"], ["内存", "4 GB", "8 GB 以上"],
        ["磁盘", "基础软件 50 MB 内", "按本地模型体积预留空间"],
        ["显卡", "基础模式不需要", "模型模式按模型规模配置 CUDA 显卡"],
    ], [2.5, 6, 7.2])
    add_heading(doc, "2.2 安装基础软件", 2)
    add_body(doc, "在命令行进入软件根目录，即包含 pyproject.toml 的目录，执行以下命令。安装过程只安装本软件基础包，默认不下载大模型。")
    add_code(doc, "python -m pip install -e .\nllmctg --version")
    add_body(doc, "看到 llmctg 1.0.0 表示入口安装成功。若不执行安装，也可以将 src 目录加入 PYTHONPATH 后使用 python -m llmctg。")
    add_heading(doc, "2.3 安装模型扩展", 2)
    add_body(doc, "只有需要加载本地 Hugging Face 模型或执行 LoRA 微调时才安装扩展依赖。模型权重不包含在本软件中，应由用户从合法来源准备。")
    add_code(doc, 'python -m pip install -e ".[llm]"')
    add_heading(doc, "3 快速演示", 1)
    add_body(doc, "快速演示不调用外部模型，会分别生成小学、初中和高中示例，完成自动评分并输出报告。")
    add_code(doc, "llmctg demo --output output/demo")
    add_table(doc, ["文件", "用途"], [
        ["articles.jsonl", "逐条保存文章、请求信息和完整评估结果"],
        ["summary.json", "保存数量、通过数、平均分和异常数"],
        ["summary.csv", "用于办公软件筛选和统计"],
        ["report.html", "用于浏览器阅读文章和质量提示"],
    ], [4, 11.7])
    add_heading(doc, "4 单篇内容生成", 1)
    add_heading(doc, "4.1 基本命令", 2)
    add_code(doc, 'llmctg generate --topic "海洋" --level junior --count 2 --seed 2026 --output output/ocean.jsonl')
    add_table(doc, ["参数", "是否必填", "说明"], [
        ["--topic", "是", "文章主题，不允许为空"], ["--level", "否", "primary、junior、senior 或中文学段，默认 junior"],
        ["--instruction", "否", "补充内容范围、体裁或表达要求"], ["--keywords", "否", "用中文或英文逗号分隔"],
        ["--count", "否", "一次生成 1 至 20 篇"], ["--seed", "否", "固定随机种子以便复现"],
        ["--backend", "否", "template 或 transformers"], ["--output", "否", "JSONL 文件；省略时打印到终端"],
    ], [3.2, 3, 9.5])
    add_heading(doc, "4.2 指定本地模型", 2)
    add_code(doc, 'llmctg generate --backend transformers --model-path D:/models/Qwen --device auto --topic "城市生态" --level senior')
    add_body(doc, "系统会在首次生成时加载模型。device 为 auto 时优先使用可用 CUDA，否则使用 CPU。为降低供应链风险，软件不会执行模型目录中的远程自定义代码。")
    add_heading(doc, "5 文本质量评估", 1)
    add_heading(doc, "5.1 评估直接输入文本", 2)
    add_code(doc, 'llmctg evaluate --text "春天来了。小草从泥土里探出头。" --level primary')
    add_heading(doc, "5.2 评估文本文件", 2)
    add_code(doc, 'llmctg evaluate --input article.txt --level junior --output output/evaluation.json')
    add_heading(doc, "5.3 理解评估结果", 2)
    add_table(doc, ["指标", "含义"], [
        ["overall_score", "结构、长度、等级匹配和流畅性的加权总分"],
        ["predicted_level", "根据句长和字词难度估算的阅读等级"],
        ["repetition_rate", "连续四字片段的重复程度"],
        ["english_ratio", "英文字符占全文字符的比例"],
        ["uncommon_character_ratio", "不在内置常用字集合中的汉字比例"],
        ["warnings", "触发阈值的质量问题"], ["suggestions", "针对问题生成的修改建议"],
    ], [5, 10.7])
    add_body(doc, "自动评估是编辑辅助工具，不能判断所有事实是否真实，也不能代替教师、编辑或学科专家的最终审核。")
    add_heading(doc, "6 批量任务处理", 1)
    add_heading(doc, "6.1 准备任务文件", 2)
    add_body(doc, "任务文件采用 JSONL，即每一行是一个独立 JSON 对象。主题 topic 必填，其他字段可省略。")
    add_code(doc, '{"topic":"森林里的水循环","level":"primary","keywords":["树叶","雨水"]}\n{"topic":"数据如何影响公共决策","level":"senior","instruction":"说明证据边界"}')
    add_heading(doc, "6.2 执行任务", 2)
    add_code(doc, "llmctg batch --input examples/tasks.jsonl --output output/batch")
    add_body(doc, "默认启用断点续作。系统为每条源任务计算稳定编号，已完成任务会跳过。若需要重新计算，可使用 --no-resume 并选择新的输出目录。单条任务失败后，原因写入 errors.jsonl，其余任务继续执行。")
    add_heading(doc, "7 训练数据准备", 1)
    add_body(doc, "原始文本每行由正文、制表符和等级标签构成。等级可使用 0、1、2 或小学、初中、高中。转换器会生成带 instruction、input、output 和 level 字段的 JSONL。")
    add_code(doc, "llmctg prepare-data --input train.txt --output data/train.jsonl")
    add_heading(doc, "8 LoRA 模型微调", 1)
    add_body(doc, "微调前应确认基础模型和训练语料的授权，准备足够磁盘与显存，并先使用少量样本验证流程。软件使用显式参数替代原始脚本中的个人绝对路径。")
    add_code(doc, "llmctg train --model-path D:/models/Qwen --dataset data/train.jsonl --output models/reading-lora --epochs 3 --batch-size 1 --max-length 1024 --device auto")
    add_bullets(doc, [
        "训练输出目录保存 LoRA 适配器和分词器配置。",
        "CUDA 环境启用半精度与梯度检查点；CPU 环境使用单精度。",
        "若模型层名称与 q_proj、k_proj、v_proj、o_proj 不匹配，应由维护人员调整目标模块。",
        "训练损失只反映拟合过程，不能代替独立测试集和人工内容评价。",
    ])
    add_heading(doc, "9 配置说明", 1)
    add_body(doc, "执行 init-config 可生成默认 JSON 配置。命令行中的后端、模型路径和设备会覆盖配置文件对应值。")
    add_code(doc, "llmctg init-config --output llmctg.json\nllmctg --config llmctg.json generate --topic \"气候\"")
    add_table(doc, ["配置项", "默认值", "说明"], [
        ["generation.backend", "template", "生成后端"], ["generation.device", "auto", "推理设备"],
        ["generation.temperature", "0.9", "采样随机性"], ["generation.top_p", "0.9", "核采样阈值"],
        ["generation.repetition_penalty", "1.1", "模型重复惩罚"], ["evaluation.pass_score", "60", "质量通过分数"],
        ["evaluation.repetition_warning", "0.18", "重复率警告阈值"], ["log_level", "INFO", "日志级别"],
    ], [6, 3, 6.7])
    add_heading(doc, "10 输出文件说明", 1)
    add_body(doc, "JSON 和 JSONL 保存全部结构化信息，适合程序继续处理；CSV 以 UTF-8 BOM 编码，适合 Excel 等软件；HTML 报告完全本地化，可直接双击打开。生成时间采用带时区的 ISO 8601 格式。")
    add_heading(doc, "11 常见问题与故障处理", 1)
    add_table(doc, ["现象", "原因与处理方法"], [
        ["提示主题不能为空", "检查 --topic 或批量任务 topic 字段，删除只有空格的值"],
        ["提示不支持的阅读等级", "使用 primary、junior、senior、0、1、2 或相应中文学段"],
        ["提示缺少模型依赖", "安装 [llm] 扩展，并确认当前命令使用同一 Python 环境"],
        ["模型目录不存在", "检查路径拼写与访问权限，不要填写网络模型名称代替本地目录"],
        ["CUDA 显存不足", "降低最大令牌数、使用更小模型或改用 CPU"],
        ["批量任务部分失败", "查看 errors.jsonl 的行号、异常类型和消息，修正后使用新目录重跑"],
        ["中文在终端显示异常", "将终端切换到 UTF-8，文件本身仍按 UTF-8 保存"],
        ["质量分低", "查看 warnings 与 suggestions，重点检查长度、长句、重复和等级匹配"],
    ], [5.2, 10.5])
    add_heading(doc, "12 数据安全与使用边界", 1)
    add_bullets(doc, [
        "默认后端完全离线，不向网络发送主题或文章。",
        "本地模型后端只读取用户指定目录，但模型、语料和数据库授权由使用者负责确认。",
        "请勿把含个人敏感信息、未授权作品或保密材料直接用于训练。",
        "生成内容投入教学、出版或公开传播前，必须进行事实、版权、年龄适宜性和价值导向审核。",
        "备份输出目录时应遵守所在单位的数据分类与保存制度。",
    ])
    add_heading(doc, "13 软件卸载", 1)
    add_body(doc, "使用 pip 安装时执行下列命令卸载程序。输出目录、模型和训练数据不会自动删除，应由用户确认后单独归档或清理。")
    add_code(doc, "python -m pip uninstall llmctg")
    DELIVERY.mkdir(parents=True, exist_ok=True)
    target = DELIVERY / f"{SOFTWARE_NAME}{VERSION}用户操作手册.docx"
    doc.save(target)
    return target


def source_files():
    include = [ROOT / "pyproject.toml"]
    for base in (ROOT / "src", ROOT / "scripts"):
        for pattern in ("*.py", "*.ps1", "*.sh"):
            include.extend(base.rglob(pattern))
    include = [path for path in include if path.name != "build_copyright_docs.py"]
    return sorted(set(include), key=lambda p: p.as_posix())


def build_source_document():
    lines = []
    for path in source_files():
        relative = path.relative_to(ROOT).as_posix()
        content = path.read_text(encoding="utf-8-sig").replace("\r\n", "\n").replace("\r", "\n")
        lines.append((relative, (None, f"===== 文件开始 {relative} =====")))
        for number, line in enumerate(content.split("\n"), 1):
            lines.append((relative, (number, line)))
    code_lines = [(path, item) for path, item in lines if item[0] is not None]
    doc = Document()
    configure_document(doc, f"{SOFTWARE_NAME} {VERSION} 源程序")
    section = doc.sections[0]
    section.top_margin = Cm(1.3)
    section.bottom_margin = Cm(1.3)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)
    page_line_count = 0
    page_number = 1
    for entry_index, (path, item) in enumerate(lines):
        number, text = item
        if page_line_count == 0:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(3)
            r = p.add_run(f"源程序第 {page_number} 页    {path}")
            set_font(r, "黑体", "Arial", 8.5, True)
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        p.paragraph_format.line_spacing = Pt(9.3)
        if number is None:
            r = p.add_run(text)
            set_font(r, "黑体", "Consolas", 6.8, True)
        else:
            r = p.add_run(f"{number:04d}  {text.expandtabs(4)}")
            set_font(r, "等线", "Consolas", 6.8)
        page_line_count += 1
        if page_line_count == 50:
            if entry_index < len(lines) - 1:
                doc.add_page_break()
            page_line_count = 0
            page_number += 1
    DELIVERY.mkdir(parents=True, exist_ok=True)
    target = DELIVERY / f"{SOFTWARE_NAME}{VERSION}源程序鉴别材料.docx"
    doc.save(target)
    return target


if __name__ == "__main__":
    manual = build_manual()
    source = build_source_document()
    print(json.dumps({"manual": str(manual), "source": str(source)}, ensure_ascii=False, indent=2))
